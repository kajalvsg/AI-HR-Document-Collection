from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document as DocxDocument

from app.utils.errors import AppError


class ResumeParserService:
    @staticmethod
    def parse(file_path: str, extension: str) -> str:
        path = Path(file_path)
        if not path.exists():
            raise AppError("Resume file not found on server", 500, "parse_error")

        ext = extension.lower()
        try:
            if ext == "pdf":
                text = ResumeParserService._parse_pdf(path)
            elif ext == "docx":
                text = ResumeParserService._parse_docx(path)
            else:
                raise AppError(
                    f"Unsupported resume format: {ext}",
                    400,
                    "invalid_file_type",
                )
        except AppError:
            raise
        except Exception as exc:
            raise AppError(
                f"Failed to parse resume: {exc}",
                422,
                "parsing_failure",
            ) from exc

        return ResumeParserService._normalize_text(text)

    @staticmethod
    def _parse_pdf(path: Path) -> str:
        reader = PdfReader(str(path))
        if not reader.pages:
            return ""
        parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
        return "\n".join(parts)

    @staticmethod
    def _parse_docx(path: Path) -> str:
        doc = DocxDocument(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    @staticmethod
    def _normalize_text(text: str) -> str:
        lines = [line.strip() for line in text.replace("\r", "\n").split("\n")]
        cleaned = "\n".join(line for line in lines if line)
        return cleaned.strip()
